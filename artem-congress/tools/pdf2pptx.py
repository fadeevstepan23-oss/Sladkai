"""PDF (из HTML-презентации) -> PPTX.

Режим editable: фон каждой страницы — растр без текста (фото, линии, планы),
поверх — редактируемый текст: одна строка = один текстовый блок, шрифт/кегль/цвет/
трекинг восстановлены из PDF.
Режим image: каждая страница — одна картинка во весь слайд (точная копия).

python3 pdf2pptx.py in.pdf out.pptx fonts_dir work_dir [editable|image]
"""
import io
import re
import subprocess
import sys

import pymupdf
from fontTools.ttLib import TTFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

SRC, OUT, FONTS, WORK = sys.argv[1:5]
MODE = sys.argv[5] if len(sys.argv) > 5 else "editable"
ZOOM = 2.0                      # 1440x810 pt -> 2880x1620 px
SLIDE_W, SLIDE_H = 12192000, 6858000

doc = pymupdf.open(SRC)
PW, PH = doc[0].rect.width, doc[0].rect.height
K = SLIDE_W / PW                # EMU на пункт PDF
FS = (SLIDE_W / 12700) / PW     # пункты PPTX на пункт PDF (960/1440)


def emu(v):
    return Emu(int(round(v * K)))


# ---------- шрифты: xref Type3 -> реальное начертание ----------
xref_name = {}
for line in subprocess.run(["pdffonts", SRC], capture_output=True, text=True).stdout.splitlines()[2:]:
    parts = line.split()
    if len(parts) >= 3:
        xref_name[int(parts[-2])] = parts[0].split("+", 1)[-1]


def face_for(name):
    """-> (typeface в PPTX, bold, ttf для метрик)"""
    if name.startswith("JetBrainsMono"):
        if "Medium" in name:
            return "JetBrains Mono Medium", False, "JetBrainsMono-Medium"
        return "JetBrains Mono", False, "JetBrainsMono-Regular"
    if name.startswith("Onest"):
        m = re.search(r"wght([0-9A-F]+)", name)
        if m:
            w = int(m.group(1), 16) / 65536
        elif "Light" in name:
            w = 300
        elif "Medium" in name:
            w = 500
        else:
            w = 400
        if w <= 250:
            return "Onest ExtraLight", False, "Onest-ExtraLight"
        if w <= 370:
            return "Onest Light", False, "Onest-Light"
        if w <= 470:
            return "Onest", False, "Onest-Regular"
        if w <= 650:
            return "Onest Medium", False, "Onest-Medium"
        return "Onest", True, "Onest-Bold"
    if "Light" in name:
        return "Arial", False, None
    return "Arial", False, None


metrics = {}


def font_metrics(ttf):
    if ttf not in metrics:
        if ttf is None:
            metrics[ttf] = None
        else:
            f = TTFont(f"{FONTS}/{ttf}.ttf")
            os2 = f["OS/2"]
            metrics[ttf] = dict(upm=f["head"].unitsPerEm, asc=os2.sTypoAscender, desc=-os2.sTypoDescender,
                                cmap=f.getBestCmap(), hmtx=f["hmtx"].metrics)
    return metrics[ttf]


def natural_width(text, ttf, size):
    m = font_metrics(ttf)
    if not m:
        return None
    total = 0
    for ch in text:
        g = m["cmap"].get(ord(ch))
        if g is None:
            g = m["cmap"].get(ord(" ")) if ch in "  " else None
        if g is None:
            return None
        total += m["hmtx"][g][0]
    return total / m["upm"] * size


def span_font(span):
    m = re.search(r"\((\d+) 0 R\)", span["font"])
    name = xref_name.get(int(m.group(1)), span["font"]) if m else span["font"]
    return face_for(name)


# ---------- сбор строк ----------
def page_lines(page):
    raw = page.get_text("rawdict")
    out = []
    for b in raw["blocks"]:
        if b["type"] != 0:
            continue
        for ln in b["lines"]:
            if abs(ln["dir"][0] - 1) > 1e-3:
                continue  # повёрнутый текст остаётся в фоне
            # PyMuPDF иногда склеивает в одну «строку» спаны с разными базовыми линиями
            # (крупная цифра рядом с подписью) — делим по базовой линии
            groups = []
            for sp in ln["spans"]:
                if not sp["chars"]:
                    continue
                if sp["chars"] and not "".join(c["c"] for c in sp["chars"]).strip():
                    if groups:
                        groups[-1].append(sp)
                    continue
                if groups and abs(groups[-1][0]["origin"][1] - sp["origin"][1]) < 2:
                    groups[-1].append(sp)
                else:
                    groups.append([sp])
            for g in groups:
                out.extend(build_runs(g, ln["bbox"]))
    return out


def build_runs(spans, bbox):
    runs = []
    for sp in spans:
        chars = sp["chars"]
        if not chars:
            continue
        face, bold, ttf = span_font(sp)
        key = (face, bold, round(sp["size"], 2), sp["color"])
        text = "".join(c["c"] for c in chars)
        x0 = chars[0]["bbox"][0]
        x1 = chars[-1]["bbox"][2]
        if runs and runs[-1]["key"] == key:
            r = runs[-1]
            r["text"] += text
            r["x1"] = x1
            r["chars"] += chars
        else:
            runs.append(dict(key=key, ttf=ttf, text=text, x0=x0, x1=x1, chars=list(chars),
                             origin=sp["origin"], size=sp["size"]))
    # пробелы по краям строки не нужны
    while runs and not runs[-1]["text"].strip():
        runs.pop()
    while runs and not runs[0]["text"].strip():
        runs.pop(0)
    if not runs:
        return []
    return [dict(runs=runs, bbox=bbox)]


def letter_spacing(run):
    """Трекинг в пунктах PDF: (измеренная ширина - естественная) / (n-1)."""
    text = run["text"]
    n = len(text)
    if n < 2:
        return 0.0
    nat = natural_width(text, run["ttf"], run["size"])
    if nat is None:
        return 0.0
    # измеряем по началам первого и последнего символа + естественная ширина последнего
    first = run["chars"][0]["origin"][0]
    last = run["chars"][-1]["origin"][0]
    last_w = natural_width(text[-1], run["ttf"], run["size"]) or 0
    nat_to_last = nat - last_w
    sp = ((last - first) - nat_to_last) / (n - 1)
    return max(-0.2 * run["size"], min(0.6 * run["size"], sp))


# ---------- сборка ----------
prs = Presentation()
prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
blank = prs.slide_layouts[6]


def add_picture_full(slide, png_bytes):
    slide.shapes.add_picture(io.BytesIO(png_bytes), 0, 0, SLIDE_W, SLIDE_H)


def render(page, ext="jpg"):
    pix = page.get_pixmap(matrix=pymupdf.Matrix(ZOOM, ZOOM), alpha=False)
    return pix.tobytes(ext, jpg_quality=90) if ext == "jpg" else pix.tobytes(ext)


for pno in range(len(doc)):
    slide = prs.slides.add_slide(blank)
    if MODE == "image":
        add_picture_full(slide, render(doc[pno]))
        continue

    lines = page_lines(doc[pno])

    # фон без горизонтального текста
    tmp = pymupdf.open()
    tmp.insert_pdf(doc, from_page=pno, to_page=pno)
    bg = tmp[0]
    for ln in lines:
        for r in ln["runs"]:
            for c in r["chars"]:
                x0, y0, x1, y1 = c["bbox"]
                bg.add_redact_annot(pymupdf.Rect(x0 + 0.3, y0 + 0.3, x1 - 0.3, y1 - 0.3), fill=False)
    bg.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_NONE,
                        graphics=pymupdf.PDF_REDACT_LINE_ART_NONE,
                        text=pymupdf.PDF_REDACT_TEXT_REMOVE)
    add_picture_full(slide, render(bg))
    pic = slide.shapes[-1]
    pic.name = "Фон слайда"

    for ln in lines:
        runs = ln["runs"]
        base = runs[0]
        m = font_metrics(base["ttf"])
        size = max(r["size"] for r in runs if r["text"].strip())
        asc = (m["asc"] / m["upm"]) if m else 0.92
        desc = (m["desc"] / m["upm"]) if m else 0.25
        baseline = base["origin"][1]
        x0 = runs[0]["x0"]
        x1 = runs[-1]["x1"]
        # у Onest и JetBrains Mono стоит USE_TYPO_METRICS: PowerPoint (Windows и Mac)
        # и LibreOffice ставят базовую линию на typo-ascender от верха блока
        top = baseline - asc * size
        h = (asc + desc) * size
        w = (x1 - x0) * 1.04 + size * 0.5
        tb = slide.shapes.add_textbox(emu(x0), emu(top), emu(w), emu(h))
        tf = tb.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.word_wrap = False
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = MSO_ANCHOR.TOP
        p = tf.paragraphs[0]
        p.line_spacing = 1.0
        for i, r in enumerate(runs):
            face, bold, rsize, color = r["key"]
            text = r["text"].replace(" ", " ")
            if i == len(runs) - 1:
                text = text.rstrip()
            if i == 0:
                text = text.lstrip()
            if not text:
                continue
            run = p.add_run()
            run.text = text
            f = run.font
            f.name = face
            f.size = Pt(round(rsize * FS * 2) / 2)
            f.bold = bold
            f.color.rgb = RGBColor((color >> 16) & 255, (color >> 8) & 255, color & 255)
            rPr = run._r.get_or_add_rPr()
            for tag in ("a:latin", "a:cs", "a:ea"):
                el = rPr.find(qn(tag))
                if el is None:
                    el = rPr.makeelement(qn(tag), {})
                    rPr.append(el)
                el.set("typeface", face)
            sp = letter_spacing(r)
            if abs(sp) > 0.02 * rsize:
                rPr.set("spc", str(int(round(sp * FS * 100))))
        tb.name = "Текст: " + "".join(r["text"] for r in runs).strip()[:40]
    print(f"slide {pno + 1}: {len(lines)} строк", flush=True)

prs.core_properties.title = doc.metadata.get("title") or "Презентация"
prs.save(OUT)
print("saved", OUT)
